"""Inspect real Excel headers before pandas can mangle duplicate names."""
from dataclasses import dataclass, field
from pathlib import Path
import re
import pandas as pd
from openpyxl import load_workbook


def normalize_column_names(columns, aliases=None):
    aliases = aliases or {}
    normalized_aliases = {re.sub('_+', '_', re.sub(r'\s+', '_', key.strip().lower())): value for key, value in aliases.items()}
    result = []
    for value in columns:
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f'Empty or non-text source column name: {value!r}')
        normalized = re.sub('_+', '_', re.sub(r'\s+', '_', value.strip().lower()))
        result.append(normalized_aliases.get(normalized, normalized))
    if len(result) != len(set(result)):
        raise ValueError(f'Duplicate column names after normalization/mapping: {result}')
    return result


def validate_required_columns(frame, required, allow_extra=False):
    names = list(frame.columns)
    if len(names) != len(set(names)):
        raise ValueError('Duplicate column names are not allowed.')
    missing = sorted(set(required) - set(names))
    extra = sorted(set(names) - set(required))
    if missing or (extra and not allow_extra):
        raise ValueError(f'Source schema mismatch. Missing={missing}; unexpected={extra}. Use an explicit source.column_aliases mapping.')
    return extra


@dataclass
class SourceValidationResult:
    status: str
    source_file: str
    messages: list[str] = field(default_factory=list)
    column_mapping: dict[str, str] = field(default_factory=dict)
    detected_types: dict[str, str] = field(default_factory=dict)
    records_read: int = 0
    data: pd.DataFrame | None = field(default=None, repr=False)

    def as_dict(self):
        return {key: getattr(self, key) for key in ['status', 'source_file', 'messages', 'column_mapping', 'detected_types', 'records_read']}


class SourceValidationError(ValueError):
    def __init__(self, result):
        self.result = result
        super().__init__('; '.join(result.messages))


def validate_source(path, config):
    path = Path(path)
    result = SourceValidationResult('FAIL', str(path))
    try:
        if not path.is_file():
            raise ValueError(f'Excel source file does not exist: {path}')
        if path.suffix.lower() not in {'.xlsx', '.xlsm'}:
            raise ValueError(f'Unsupported Excel extension {path.suffix!r}; use .xlsx or .xlsm.')
        source = config['source']
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = source['sheet_name'] or workbook.sheetnames[0]
            if sheet not in workbook.sheetnames:
                raise ValueError(f'Expected sheet {sheet!r} not found. Available sheets: {workbook.sheetnames}')
            header = list(next(workbook[sheet].iter_rows(min_row=source['header_row'] + 1, max_row=source['header_row'] + 1, values_only=True), ()))
            while header and header[-1] is None:
                header.pop()
            names = normalize_column_names(header, source.get('column_aliases'))
        finally:
            workbook.close()
        data = pd.read_excel(path, sheet_name=sheet, header=source['header_row'], dtype=object, keep_default_na=False, na_values=[''])
        if len(data.columns) != len(names):
            raise ValueError('Column count differs from nonempty headers; check blank headers or stray worksheet data.')
        data.columns = names
        reserved = {'_batch_id', '_ingestion_timestamp', '_ingestion_date', '_source_file', '_source_sha256', '_source_row_number'}
        if reserved & set(names):
            raise ValueError(f'Source columns collide with reserved ingestion metadata: {sorted(reserved & set(names))}')
        extra = validate_required_columns(data, source['required_columns'], source.get('allow_extra_columns', False))
        if data.empty or not data.astype('string').apply(lambda column: column.str.strip().ne('').fillna(False)).any(axis=None):
            raise ValueError('Excel source dataset is completely empty.')
        result.column_mapping = dict(zip(header, names))
        changed = {old: new for old, new in result.column_mapping.items() if old != new}
        if changed:
            result.messages.append(f'Explicit/syntactic column normalization: {changed}')
        if extra:
            result.messages.append(f'Configured extra columns preserved in Bronze: {extra}')
        result.status = 'WARNING' if result.messages else 'PASS'
        result.records_read = len(data)
        result.detected_types = {col: str(data[col].infer_objects().dtype) for col in data}
        result.data = data
        return result
    except Exception as error:
        result.messages.append(f'Source validation failed ({type(error).__name__}): {error}')
        raise SourceValidationError(result) from error
